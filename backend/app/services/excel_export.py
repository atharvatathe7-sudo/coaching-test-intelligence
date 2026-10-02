"""
Excel (.xlsx) export of one test.

This module is an output layer only. Every figure comes from the
existing authoritative calculations:

  analyze_batch        overview, students, chapters, topics, flagged questions
  analyze_questions    one row per question
  get_action_outcomes  observed change for each teacher action
  compare_tests        progress against the previous test

The only logic added here is presentation and the selection of the
previous test, which mirrors the dashboard's "Progress vs Previous Test"
rule (latest earlier test, same batch and subject).

All user-controlled text goes through safe_cell_value(), so ordinary
text can never be interpreted as an Excel formula.
"""

import io
import math
import re
import unicodedata
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..database.models import (
    Batch,
    Chapter,
    Institute,
    Question,
    Student,
    StudentAnswer,
    TeacherAction,
    Test,
    Topic,
)
from .action_outcomes import get_action_outcomes
from .analytics import analyze_batch, analyze_questions
from .progress import compare_tests

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

SHEET_NAMES = (
    "Test Summary",
    "Student Results",
    "Question Analysis",
    "Chapter Analysis",
    "Topic Analysis",
    "Answer Matrix",
    "Teacher Actions",
    "Test Comparison",
)

# ---------------------------------------------------------------------
# Safe cell values
# ---------------------------------------------------------------------

# A string starting with one of these can be run as a formula by Excel
# (and by LibreOffice/Sheets). Tab and carriage return are included
# because some spreadsheet programs skip them before checking.
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")

# Control characters XML cannot carry; openpyxl refuses them.
_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

MAX_CELL_CHARS = 32767


def safe_cell_value(value):
    """
    The one place user-controlled content is made safe for a cell.

    Numbers, booleans, dates and None pass through untouched, so real
    numeric values are never turned into text. Strings lose characters
    XML cannot hold, are cut to Excel's cell limit, and are prefixed with
    an apostrophe when they start with a formula trigger so they stay
    plain text.
    """

    if value is None or isinstance(value, bool):
        return value

    if isinstance(value, (int, date, datetime)):
        return value

    if isinstance(value, float):
        return value if math.isfinite(value) else None

    text = _ILLEGAL_XML.sub("", str(value))

    if len(text) > MAX_CELL_CHARS - 1:
        text = text[: MAX_CELL_CHARS - 2] + "…"

    if text.startswith(FORMULA_TRIGGERS):
        text = "'" + text

    return text


def export_filename(test: Test) -> str:
    """Deterministic, filesystem-safe name: <test name>_<date>.xlsx."""

    folded = (
        unicodedata.normalize("NFKD", test.name or "")
        .encode("ascii", "ignore")
        .decode("ascii")
    )
    stem = re.sub(r"[^A-Za-z0-9]+", "_", folded).strip("_")[:80] or "Test"

    return f"{stem}_{test.test_date.isoformat()}.xlsx"


# ---------------------------------------------------------------------
# Writing helpers
# ---------------------------------------------------------------------

PCT = "0.0%"
DECIMAL = "0.00"
DATE = "yyyy-mm-dd"
DATETIME = "yyyy-mm-dd hh:mm"
# Differences between two percentages are percentage points, not percent.
PP = '+0.00" pp";-0.00" pp";0.00" pp"'
SIGNED = "+0.00;-0.00;0.00"

HEADER_FONT = Font(bold=True)
NOTE_FONT = Font(italic=True, color="555555")
WRAP = Alignment(wrap_text=True, vertical="top")

NO_QUESTIONS = "This test has no questions, so no analysis is available."
NO_ANSWERS = (
    "No student answers have been imported for this test, "
    "so no analysis is available."
)
UNAVAILABLE = "The analysis for this test is not available."


def _pct(value):
    """Percentages are stored as 0-100 by the analytics; Excel wants 0-1."""
    return None if value is None else round(value / 100, 6)


def _as_date(value):
    """The analytics return dates as date objects or ISO strings."""
    if value is None or isinstance(value, date):
        return value
    return date.fromisoformat(value)


def _put(ws, row, column, value, number_format=None, font=None, wrap=False):
    cell = ws.cell(row=row, column=column, value=safe_cell_value(value))

    if isinstance(cell.value, str):
        cell.data_type = "s"

    if number_format:
        cell.number_format = number_format

    if font:
        cell.font = font

    if wrap:
        cell.alignment = WRAP

    return cell


def _fit_columns(ws, wrapped=(), minimum=8, maximum=60):
    for index, column in enumerate(ws.columns, start=1):
        longest = max(
            (len(str(c.value)) for c in column if c.value is not None),
            default=0,
        )
        width = max(minimum, min(longest + 2, 80 if index in wrapped else maximum))
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_table(
    ws,
    headers,
    rows,
    formats=None,
    wrapped=(),
    empty_message=None,
    freeze="A2",
):
    """
    Header row + data rows on a fresh sheet: bold header, frozen panes,
    autofilter. `formats` maps a 1-based column to a number format;
    `wrapped` lists 1-based columns with wrapped text.
    """

    formats = formats or {}

    for column, header in enumerate(headers, start=1):
        _put(ws, 1, column, header, font=HEADER_FONT, wrap=True)

    for row_number, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            _put(
                ws,
                row_number,
                column,
                value,
                number_format=formats.get(column),
                wrap=column in wrapped,
            )

    if rows:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        )
    elif empty_message:
        _put(ws, 2, 1, empty_message, font=NOTE_FONT)

    ws.freeze_panes = freeze
    _fit_columns(ws, wrapped=wrapped)


# ---------------------------------------------------------------------
# Gathering (queries and calls to the existing analytics only)
# ---------------------------------------------------------------------

def _try(function, *args):
    try:
        return function(*args)
    except ValueError:
        return None


def previous_test_for(db: Session, test: Test) -> Test | None:
    """
    Latest earlier test for the same batch and subject: the rule the
    dashboard uses for "Progress vs Previous Test".
    """

    return (
        db.query(Test)
        .filter(
            Test.id != test.id,
            Test.batch_id == test.batch_id,
            Test.subject == test.subject,
            (Test.test_date < test.test_date)
            | ((Test.test_date == test.test_date) & (Test.id < test.id)),
        )
        .order_by(Test.test_date.desc(), Test.id.desc())
        .first()
    )


def _names(db: Session, model, ids) -> dict:
    ids = {i for i in ids if i is not None}

    if not ids:
        return {}

    return {
        row.id: row.name
        for row in db.query(model).filter(model.id.in_(ids))
    }


def gather(db: Session, test: Test) -> dict:
    batch = db.get(Batch, test.batch_id)
    institute = db.get(Institute, batch.institute_id)

    question_count = (
        db.query(Question).filter(Question.test_id == test.id).count()
    )
    answer_count = (
        db.query(StudentAnswer)
        .filter(StudentAnswer.test_id == test.id)
        .count()
    )

    data = {
        "test": test,
        "batch": batch,
        "institute": institute,
        "total_students": (
            db.query(Student).filter(Student.batch_id == batch.id).count()
        ),
        "question_count": question_count,
        "notice": None,
        "questions": [],
        "batch_analysis": None,
    }

    if question_count == 0:
        data["notice"] = NO_QUESTIONS
    else:
        data["questions"] = _try(analyze_questions, db, test.id) or []

        if answer_count == 0:
            data["notice"] = NO_ANSWERS
        else:
            data["batch_analysis"] = _try(analyze_batch, db, test.id)

            if data["batch_analysis"] is None:
                data["notice"] = UNAVAILABLE

    data["question_chapters"] = _names(
        db, Chapter, [q["chapter_id"] for q in data["questions"]]
    )
    data["question_topics"] = _names(
        db, Topic, [q["topic_id"] for q in data["questions"]]
    )

    data["answers"] = {
        (row.student_id, row.question_id): row.answer
        for row in db.query(StudentAnswer).filter(
            StudentAnswer.test_id == test.id
        )
    }

    data["actions"] = (
        db.query(TeacherAction)
        .filter(TeacherAction.test_id == test.id)
        .order_by(TeacherAction.created_at, TeacherAction.id)
        .all()
    )
    data["outcomes"] = {
        outcome["action"]["id"]: outcome
        for outcome in (_try(get_action_outcomes, db, test.id) or [])
    }
    data["action_questions"] = {
        q.id: q.question_number
        for q in db.query(Question).filter(
            Question.id.in_(
                {a.question_id for a in data["actions"] if a.question_id}
            )
        )
    } if data["actions"] else {}
    data["action_chapters"] = _names(
        db, Chapter, [a.chapter_id for a in data["actions"]]
    )
    data["action_topics"] = _names(
        db, Topic, [a.topic_id for a in data["actions"]]
    )

    data["previous_test"] = previous_test_for(db, test)
    data["comparison"] = (
        _try(compare_tests, db, data["previous_test"].id, test.id)
        if data["previous_test"]
        else None
    )

    return data


# ---------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------

def _summary_sheet(ws, data):
    test, batch, institute = data["test"], data["batch"], data["institute"]
    analysis = data["batch_analysis"]
    overview = analysis["overview"] if analysis else None

    rows = [
        ("Institute", institute.name, None),
        ("Batch", batch.name, None),
        ("Test", test.name, None),
        ("Test date", test.test_date, DATE),
        ("Subject", test.subject, None),
        ("Marks per correct answer", test.marks_correct, None),
        ("Marks per wrong answer", test.marks_wrong, None),
        ("Marks per unattempted question", test.marks_blank, None),
        ("Total students in batch", data["total_students"], "0"),
        (
            "Students evaluated",
            analysis["student_count"] if analysis else 0,
            "0",
        ),
        ("Total questions", data["question_count"], "0"),
    ]

    if overview:
        rows += [
            ("Average marks", overview["average_marks"], DECIMAL),
            ("Highest marks", overview["highest_marks"], None),
            ("Lowest marks", overview["lowest_marks"], None),
            (
                "Average accuracy (mean of student accuracies)",
                _pct(overview["average_accuracy"]),
                PCT,
            ),
            (
                "Batch accuracy (all attempted answers together)",
                _pct(overview["batch_accuracy"]),
                PCT,
            ),
            ("Total correct answers", overview["total_correct"], "0"),
            ("Total wrong answers", overview["total_wrong"], "0"),
            ("Total unattempted answers", overview["total_blank"], "0"),
        ]

    for index, (label, value, number_format) in enumerate(rows, start=1):
        _put(ws, index, 1, label, font=HEADER_FONT)
        _put(ws, index, 2, value, number_format=number_format)

    row = len(rows) + 2

    if data["notice"]:
        _put(ws, row, 1, data["notice"], font=NOTE_FONT)
        row += 2

    if analysis:
        # Chapters and questions as already identified by the application.
        ranked = sorted(
            (c for c in analysis["chapters"] if c["responses"] > 0),
            key=lambda c: (c["correct_percentage"], c["chapter_name"]),
        )
        levels = {c["correct_percentage"] for c in ranked}

        if len(levels) == 1 and len(ranked) > 1:
            # A ranking of equal values would be arbitrary.
            _put(
                ws,
                row,
                1,
                "All chapters have the same correct %",
                font=HEADER_FONT,
            )
            _put(ws, row, 2, _pct(ranked[0]["correct_percentage"]), PCT)
            row += 2
        else:
            # Never list a chapter as both lowest and highest, and never
            # label a chapter lowest/highest when it ties with the other end.
            count = min(3, len(ranked) // 2)
            lowest = [
                c for c in ranked[:count]
                if c["correct_percentage"] < max(levels)
            ]
            highest = [
                c for c in reversed(ranked[-count:])
                if c["correct_percentage"] > min(levels)
            ] if count else []

            for title, chapters in (
                ("Lowest correct % chapters", lowest),
                ("Highest correct % chapters", highest),
            ):
                if not chapters:
                    continue
                _put(ws, row, 1, title, font=HEADER_FONT)
                row += 1
                for chapter in chapters:
                    _put(ws, row, 1, chapter["chapter_name"])
                    _put(ws, row, 2, _pct(chapter["correct_percentage"]), PCT)
                    row += 1
                row += 1

        flagged = analysis["difficult_questions"]
        if flagged:
            _put(
                ws, row, 1, "Questions flagged as difficult", font=HEADER_FONT
            )
            row += 1
            for question in flagged:
                _put(ws, row, 1, f"Question {question['question_number']}")
                _put(ws, row, 2, _pct(question["correct_percentage"]), PCT)
                row += 1
            row += 1

    notes = (
        "How to read this workbook",
        "Accuracy = correct answers / attempted questions (unattempted excluded).",
        "Correct % = correct answers / all responses (unattempted included).",
        "Figures describe this test only. They do not explain why results "
        "were as they were.",
    )
    for index, note in enumerate(notes):
        _put(ws, row + index, 1, note, font=HEADER_FONT if index == 0 else NOTE_FONT)

    _put(
        ws,
        row + len(notes) + 1,
        1,
        "Report generated (UTC)",
        font=HEADER_FONT,
    )
    _put(ws, row + len(notes) + 1, 2, datetime.utcnow(), DATETIME)

    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 36
    ws.freeze_panes = "A2"


def _topic_text(topics, limit=3):
    return "; ".join(
        f"{t['topic_name']} ({t['correct']}/{t['questions']})"
        for t in topics[:limit]
    )


def _student_sheet(ws, data):
    analysis = data["batch_analysis"]
    students = sorted(
        analysis["students"] if analysis else [],
        key=lambda s: (s["roll_number"], s["student_name"]),
    )

    rows = [
        (
            s["roll_number"],
            s["student_name"],
            s["marks"],
            _pct(s["accuracy"]),
            s["correct"],
            s["wrong"],
            s["blank"],
            s["attempted"],
            s["negative_marks"],
            _topic_text(s["strongest_topics"]),
            _topic_text(s["weakest_topics"]),
        )
        for s in students
    ]

    _write_table(
        ws,
        [
            "Roll number",
            "Student name",
            "Marks",
            "Accuracy",
            "Correct",
            "Wrong",
            "Unattempted",
            "Attempted",
            "Negative marks",
            "Highest-scoring topics, this test (correct/questions)",
            "Lowest-scoring topics, this test (correct/questions)",
        ],
        rows,
        formats={4: PCT, 5: "0", 6: "0", 7: "0", 8: "0"},
        wrapped=(10, 11),
        empty_message=data["notice"] or UNAVAILABLE,
        freeze="C2",
    )


def _question_sheet(ws, data):
    analysis = data["batch_analysis"]
    flags = {}

    if analysis:
        for key, label in (
            ("difficult_questions", "Difficult"),
            ("high_wrong_questions", "High wrong"),
            ("blank_heavy_questions", "Blank-heavy"),
        ):
            for question in analysis[key]:
                flags.setdefault(question["question_number"], []).append(label)

    rows = [
        (
            q["question_number"],
            data["question_chapters"].get(q["chapter_id"]),
            data["question_topics"].get(q["topic_id"]),
            q["correct_answer"],
            q["correct_count"],
            q["wrong_count"],
            q["blank_count"],
            q["total_responses"],
            _pct(q["correct_percentage"]),
            _pct(q["wrong_percentage"]),
            _pct(q["blank_percentage"]),
            ", ".join(flags.get(q["question_number"], [])),
        )
        for q in sorted(data["questions"], key=lambda q: q["question_number"])
    ]

    _write_table(
        ws,
        [
            "Question",
            "Chapter",
            "Topic",
            "Correct answer",
            "Correct",
            "Wrong",
            "Unattempted",
            "Total responses",
            "Correct %",
            "Wrong %",
            "Unattempted %",
            "Flags (as identified by the application)",
        ],
        rows,
        formats={
            1: "0", 5: "0", 6: "0", 7: "0", 8: "0",
            9: PCT, 10: PCT, 11: PCT,
        },
        empty_message=data["notice"] or NO_QUESTIONS,
    )


def _chapter_sheet(ws, data):
    analysis = data["batch_analysis"]
    chapters = sorted(
        analysis["chapters"] if analysis else [],
        key=lambda c: c["chapter_name"],
    )
    rows = [
        (
            c["chapter_name"],
            c["subject"],
            c["questions"],
            c["responses"],
            c["correct"],
            c["wrong"],
            c["blank"],
            _pct(c["correct_percentage"]),
            _pct(c["wrong_percentage"]),
            _pct(c["blank_percentage"]),
        )
        for c in chapters
    ]
    _write_table(
        ws,
        [
            "Chapter", "Subject", "Questions", "Total responses", "Correct",
            "Wrong", "Unattempted", "Correct %", "Wrong %", "Unattempted %",
        ],
        rows,
        formats={
            3: "0", 4: "0", 5: "0", 6: "0", 7: "0",
            8: PCT, 9: PCT, 10: PCT,
        },
        empty_message=data["notice"] or UNAVAILABLE,
    )


def _topic_sheet(ws, data):
    analysis = data["batch_analysis"]
    topics = sorted(
        analysis["topics"] if analysis else [],
        key=lambda t: (t["chapter_name"], t["topic_name"]),
    )
    rows = [
        (
            t["chapter_name"],
            t["topic_name"],
            t["questions"],
            t["responses"],
            t["correct"],
            t["wrong"],
            t["blank"],
            _pct(t["correct_percentage"]),
            _pct(t["wrong_percentage"]),
            _pct(t["blank_percentage"]),
        )
        for t in topics
    ]
    _write_table(
        ws,
        [
            "Chapter", "Topic", "Questions", "Total responses", "Correct",
            "Wrong", "Unattempted", "Correct %", "Wrong %", "Unattempted %",
        ],
        rows,
        formats={
            3: "0", 4: "0", 5: "0", 6: "0", 7: "0",
            8: PCT, 9: PCT, 10: PCT,
        },
        empty_message=data["notice"] or UNAVAILABLE,
    )


def _matrix_sheet(ws, data):
    analysis = data["batch_analysis"]
    students = sorted(
        analysis["students"] if analysis else [],
        key=lambda s: (s["roll_number"], s["student_name"]),
    )
    questions = sorted(data["questions"], key=lambda q: q["question_number"])

    headers = ["Roll number", "Student name"] + [
        f"Q{q['question_number']}" for q in questions
    ]
    rows = []

    for student in students:
        answers = []
        for question in questions:
            answer = data["answers"].get(
                (student["student_id"], question["question_id"])
            )
            answers.append((answer or "").strip() or None)
        rows.append([student["roll_number"], student["student_name"], *answers])

    _write_table(
        ws,
        headers,
        rows,
        empty_message=data["notice"] or UNAVAILABLE,
        freeze="C2",
    )

    for index in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 6
        for row in ws.iter_rows(min_row=2, min_col=index, max_col=index):
            row[0].alignment = Alignment(horizontal="center")


OUTCOME_STATUS = {
    "measured": "Measured",
    "no_target": "Not linked to a topic or chapter",
    "no_subsequent_test": "No later test with answers yet",
    "topic_absent": "Topic not in a later test yet",
    "chapter_absent": "Chapter not in a later test yet",
    "insufficient_data": "Too few responses to report a change",
}

# Same wording as the dashboard (frontend/src/ActionOutcome.jsx).
OUTCOME_CAVEATS = {
    "action_recorded_after_next_test":
        "The action was recorded after that test's date.",
    "action_still_planned": "The action is still planned.",
    "marking_scheme_differs": "Marking scheme differs between the tests.",
    "test_reevaluated_after_action":
        "This test's answer key or answers were corrected after the "
        "action was recorded; the finding shown at that time may differ.",
}


def _action_sheet(ws, data):
    rows = []

    for action in data["actions"]:
        outcome = data["outcomes"].get(action.id) or {}
        target = outcome.get("target") or {}
        next_test = outcome.get("next_test") or {}
        measured = outcome.get("status") == "measured"

        rows.append((
            action.id,
            data["action_questions"].get(action.question_id),
            data["action_chapters"].get(action.chapter_id),
            data["action_topics"].get(action.topic_id),
            action.finding_type,
            action.finding_title,
            action.finding_reason,
            action.action_type,
            action.status,
            action.note,
            action.created_at,
            action.updated_at,
            (
                f"{target['level'].capitalize()}: {target['name']}"
                if target
                else None
            ),
            next_test.get("name"),
            _as_date(next_test.get("date")),
            _pct(outcome.get("previous_percentage")) if measured else None,
            _pct(outcome.get("next_percentage")) if measured else None,
            outcome.get("change_pp") if measured else None,
            OUTCOME_STATUS.get(outcome.get("status"), outcome.get("status")),
            "; ".join(
                OUTCOME_CAVEATS.get(code, code)
                for code in outcome.get("caveats", [])
            ),
            outcome.get("note"),
        ))

    _write_table(
        ws,
        [
            "Action ID", "Question", "Chapter", "Topic", "Finding type",
            "Finding", "Finding detail", "Action type", "Status", "Note",
            "Created (UTC)", "Last updated (UTC)", "Measured target",
            "Next comparable test", "Next test date",
            "Correct % in this test", "Correct % in next test",
            "Observed change (percentage points)", "Outcome", "Caveats",
            "Interpretation",
        ],
        rows,
        formats={
            1: "0", 2: "0", 11: DATETIME, 12: DATETIME, 15: DATE,
            16: PCT, 17: PCT, 18: SIGNED,
        },
        wrapped=(6, 7, 10, 20, 21),
        empty_message="No teacher actions have been recorded for this test.",
        freeze="B2",
    )


def _comparison_sheet(ws, data):
    comparison = data["comparison"]
    previous = data["previous_test"]

    if comparison is None:
        _put(ws, 1, 1, "No comparable test is available.", font=HEADER_FONT)
        if previous is not None:
            _put(
                ws,
                2,
                1,
                f"An earlier test exists ({previous.name}) but it could not "
                "be compared, for example because it has no imported answers.",
                font=NOTE_FONT,
            )
        elif data["notice"]:
            _put(ws, 2, 1, data["notice"], font=NOTE_FONT)
        ws.column_dimensions["A"].width = 90
        return

    row = 1
    prev, cur = comparison["previous_test"], comparison["current_test"]

    for label, value, number_format in (
        ("Previous test", prev["name"], None),
        ("Previous test date", _as_date(prev["test_date"]), DATE),
        ("This test", cur["name"], None),
        ("This test date", _as_date(cur["test_date"]), DATE),
        ("Students in both tests", comparison["common_student_count"], "0"),
    ):
        _put(ws, row, 1, label, font=HEADER_FONT)
        _put(ws, row, 2, value, number_format)
        row += 1

    for note in comparison["comparability_notes"]:
        _put(ws, row, 1, "Note", font=HEADER_FONT)
        _put(ws, row, 2, note, wrap=True)
        row += 1

    def section(title, headers, body, formats):
        nonlocal row
        row += 1
        _put(ws, row, 1, title, font=HEADER_FONT)
        row += 1
        for column, header in enumerate(headers, start=1):
            _put(ws, row, column, header, font=HEADER_FONT, wrap=True)
        for values in body:
            row += 1
            for column, value in enumerate(values, start=1):
                _put(ws, row, column, value, formats.get(column))
        row += 1

    b = comparison["batch_comparison"]
    section(
        "Batch",
        ["Measure", "Previous test", "This test", "Change"],
        [
            (
                "Average marks",
                b["average_previous_marks"],
                b["average_current_marks"],
                b["average_marks_change"],
            ),
            (
                "Average accuracy",
                _pct(b["average_previous_accuracy"]),
                _pct(b["average_current_accuracy"]),
                b["average_accuracy_change"],
            ),
            (
                "Marks as % of maximum",
                _pct(b["previous_marks_percentage"]),
                _pct(b["current_marks_percentage"]),
                b["marks_percentage_change_pp"],
            ),
        ],
        {},
    )
    # Per-row formats for the batch table (mixed units).
    first = row - 3
    for offset, formats in enumerate(
        ((DECIMAL, DECIMAL, SIGNED), (PCT, PCT, PP), (PCT, PCT, PP))
    ):
        for column, fmt in zip((2, 3, 4), formats):
            ws.cell(row=first + offset, column=column).number_format = fmt

    summary = comparison["student_summary"]
    section(
        "Students (improved / declined / unchanged, by marks)",
        ["Improved", "Declined", "Unchanged"],
        [(summary["improved"], summary["declined"], summary["unchanged"])],
        {1: "0", 2: "0", 3: "0"},
    )

    for title, key, name_key, extra in (
        ("Chapter progress", "chapter_progress", "chapter_name", None),
        ("Topic progress", "topic_progress", "topic_name", "chapter_name"),
    ):
        body = []
        for item in sorted(
            comparison[key],
            key=lambda i: (i.get("chapter_name") or "", i[name_key]),
        ):
            body.append((
                item.get("chapter_name") if extra else None,
                item[name_key],
                _pct(item["previous_percentage"]),
                _pct(item["current_percentage"]),
                item["change_pp"],
                item["status"].replace("_", " "),
            ))
        section(
            f"{title} (correct %)",
            ["Chapter", "Name", "Previous test", "This test",
             "Change (percentage points)", "Status"],
            body,
            {3: PCT, 4: PCT, 5: SIGNED},
        )

    section(
        "Student progress",
        ["Roll number", "Student name", "Marks (previous)", "Marks (this)",
         "Marks change", "Accuracy (previous)", "Accuracy (this)",
         "Accuracy change (percentage points)"],
        [
            (
                s["roll_number"], s["student_name"], s["previous_marks"],
                s["current_marks"], s["marks_change"],
                _pct(s["previous_accuracy"]), _pct(s["current_accuracy"]),
                s["accuracy_change"],
            )
            for s in sorted(
                comparison["students"], key=lambda s: s["roll_number"]
            )
        ],
        {6: PCT, 7: PCT, 8: PP},
    )

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 34
    for letter in "CDEFGH":
        ws.column_dimensions[letter].width = 18
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------

def build_test_workbook(db: Session, test: Test) -> bytes:
    """The finished .xlsx for one test, in memory."""

    data = gather(db, test)
    workbook = Workbook()
    workbook.remove(workbook.active)

    builders = (
        _summary_sheet,
        _student_sheet,
        _question_sheet,
        _chapter_sheet,
        _topic_sheet,
        _matrix_sheet,
        _action_sheet,
        _comparison_sheet,
    )

    for name, build in zip(SHEET_NAMES, builders):
        build(workbook.create_sheet(name), data)

    buffer = io.BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()
