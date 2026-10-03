"""
OMR import through the API (Stage 2B), running the real recogniser.

Synthetic sheets with known content are drawn by tests/omr_sheets.py and
sent to POST /api/tests/{id}/omr/import, then the result is compared with
numbers calculated by hand below, not by calling production evaluation
code.

Test "OMR Test" (10 questions, marking +4 / -1 / 0), answer key:

    Q1 A  Q2 B  Q3 C  Q4 D  Q5 A  Q6 B  Q7 C  Q8 D  Q9 A  Q10 B

Students and what they marked:

  101001  A B C D A B C D A B           10 correct                -> 40
  101002  A C C - A B D D - B           correct Q1 Q3 Q5 Q6 Q8 Q10 (6)
                                        wrong   Q2 Q7 (2)
                                        blank   Q4 Q9 (2)        -> 24 - 2 = 22
  101003  (nothing marked)              10 blank                  -> 0
  101004  B B A D B A C A A A           correct Q2 Q4 Q7 Q9 (4)
                                        wrong   Q1 Q3 Q5 Q6 Q8 Q10 (6)
                                                                 -> 16 - 6 = 10
  101005  A B C D A B C D A B           10 correct                -> 40
          (in the "ambiguous" run Q1 is marked "AB" instead of "A")
"""

import json
import re
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

import io

from backend.app import config
from backend.app.database import models
from backend.app.database.connection import SessionLocal
from backend.app.omr import checker

import manage
from conftest import make_client
from omr_sheets import render_sheet

PASSWORD = "omr-test-password"
KEY = "ABCDABCDAB"

STUDENTS = {
    "101001": "ABCDABCDAB",
    "101002": "ACC-ABDD-B",
    "101003": "----------",
    "101004": "BBADBACAAA",
    "101005": "ABCDABCDAB",
}
EXPECTED = {            # roll: (marks, correct, wrong, blank)
    "101001": (40, 10, 0, 0),
    "101002": (22, 6, 2, 2),
    "101003": (0, 0, 0, 10),
    "101004": (10, 4, 6, 0),
    "101005": (40, 10, 0, 0),
}

_counter = {"n": 0}


def upload(client, path, text, data):
    response = client.post(
        f"/api/imports/{path}",
        data={k: str(v) for k, v in data.items()},
        files={"file": ("f.csv", text.encode(), "text/csv")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def make_test(ctx, name=None, questions=10, key=KEY):
    """A new test in the OMR batch with its answer key; returns its id."""
    _counter["n"] += 1
    key_csv = "question_number,correct_answer,chapter,topic\n" + "".join(
        f"{q},{key[(q - 1) % len(key)]},{'Mechanics' if q <= 5 else 'Optics'},"
        f"{'Kinematics' if q <= 5 else 'Lenses'}\n"
        for q in range(1, questions + 1)
    )
    result = upload(
        ctx["admin"], "test-setup", key_csv,
        dict(batch_id=ctx["batch_id"], test_name=name or f"OMR Test {_counter['n']}",
             subject="Physics", test_date="2026-10-02", dry_run=False,
             confirm_new_chapters_topics=True),
    )
    assert result["status"] == "imported", result
    return result["summary"]["test_id"]


@pytest.fixture(scope="module")
def ctx(demo_database):
    db = SessionLocal()
    try:
        institute = models.Institute(name="OMR Institute")
        db.add(institute)
        db.flush()
        manage.make_user(db, institute.id, "OMR Admin", "admin@omr.test", "admin", PASSWORD)
        manage.make_user(db, institute.id, "OMR Teacher", "teacher@omr.test", "teacher", PASSWORD)
        db.commit()
    finally:
        db.close()

    admin = make_client("admin@omr.test", PASSWORD)
    teacher = make_client("teacher@omr.test", PASSWORD)
    batch = admin.post("/api/batches", json={"name": "OMR Batch"}).json()

    names = {"101001": "Asha", "101002": "Ben", "101003": "Chitra",
             "101004": "=EVIL()+1", "101005": "Zoë 日本語"}
    roster = "roll_number,name\n" + "".join(f'{r},"{n}"\n' for r, n in names.items())
    assert upload(admin, "roster", roster, dict(batch_id=batch["id"], dry_run=False))["status"] == "imported"

    return {"admin": admin, "teacher": teacher, "batch_id": batch["id"]}


def sheets_for(rolls=None, overrides=None, names=None):
    """[(filename, png bytes)] for the given students."""
    files = []
    for i, roll in enumerate(rolls or STUDENTS, start=1):
        letters = (overrides or {}).get(roll, STUDENTS[roll])
        answers = {q: ("" if ch == "-" else ch) for q, ch in enumerate(letters, start=1)} \
            if isinstance(letters, str) else letters
        name = (names or {}).get(roll, f"scan_{roll}.png")
        files.append((name, render_sheet(roll, answers, seed=i)))
    return files


def post_omr(client, test_id, files, **form):
    response = client.post(
        f"/api/tests/{test_id}/omr/import",
        files=[("files", (name, data, "image/png")) for name, data in files],
        data={k: str(v) for k, v in form.items()},
    )
    return response


def answer_count(test_id):
    db = SessionLocal()
    try:
        return db.query(models.StudentAnswer).filter_by(test_id=test_id).count()
    finally:
        db.close()


def student_results(client, test_id):
    rows = client.get(f"/api/tests/{test_id}/analytics/students").json()["students"]
    return {r["roll_number"]: (r["marks"], r["correct"], r["wrong"], r["blank"]) for r in rows}


# ------------------------------------------------------------------
# Access control
# ------------------------------------------------------------------

def test_unauthenticated_access_is_rejected(ctx):
    test_id = make_test(ctx)
    anonymous = make_client()
    response = post_omr(anonymous, test_id, sheets_for(["101001"]))
    assert response.status_code == 401


def test_teachers_cannot_import_omr(ctx):
    test_id = make_test(ctx)
    assert post_omr(ctx["teacher"], test_id, sheets_for(["101001"])).status_code == 403


def test_other_institutes_cannot_use_this_test(ctx, client):
    """The demo institute's admin is told the test does not exist."""
    test_id = make_test(ctx)
    foreign = post_omr(client, test_id, sheets_for(["101001"]))
    missing = post_omr(client, 99999999, sheets_for(["101001"]))

    assert foreign.status_code == missing.status_code == 404
    assert foreign.json() == missing.json()
    assert answer_count(test_id) == 0


def test_demo_test_cannot_be_reached_from_this_institute(ctx, test_ids):
    response = post_omr(ctx["admin"], test_ids[0], sheets_for(["101001"]))
    assert response.status_code == 404


# ------------------------------------------------------------------
# End to end against hand-calculated ground truth
# ------------------------------------------------------------------

def test_dry_run_checks_everything_and_writes_nothing(ctx):
    test_id = make_test(ctx)
    response = post_omr(ctx["admin"], test_id, sheets_for(overrides={"101005": "ABCDABCDAB"}))
    body = response.json()

    assert response.status_code == 200, body
    assert body["status"] == "accepted" and body["committed"] is False
    assert body["counts"]["processed"] == 5 and body["counts"]["accepted"] == 5
    assert answer_count(test_id) == 0


def test_end_to_end_matches_independently_calculated_results(ctx):
    test_id = make_test(ctx)
    response = post_omr(ctx["admin"], test_id, sheets_for(), dry_run=False)
    body = response.json()

    assert response.status_code == 200, body
    assert body["status"] == "imported" and body["committed"] is True
    assert body["errors"] == []
    # 5 students x 10 questions: blanks are 2 (101002) + 10 (101003) = 12
    assert body["counts"] == {
        "submitted": 5, "processed": 5, "accepted": 5, "unreadable": 0,
        "duplicate": 0, "unknown_student": 0, "recognized": 38, "blank": 12,
        "multi_mark": 0, "invalid": 0, "review_required": 0,
    }
    assert body["import"]["students_evaluated"] == 5

    # identity and marks of every student, from the existing evaluation
    assert student_results(ctx["admin"], test_id) == EXPECTED

    # every recorded answer, read back through Question Investigation,
    # must be exactly what was drawn on the sheet (None = left blank)
    for question in range(1, 11):
        rows = ctx["admin"].get(
            f"/api/tests/{test_id}/questions/{question}/investigation"
        ).json()["students"]
        recorded = {r["roll_number"]: r["answer"] for r in rows}
        for roll, letters in STUDENTS.items():
            drawn = letters[question - 1]
            assert recorded[roll] == (None if drawn == "-" else drawn), (roll, question)


def test_multi_mark_is_detected_distinct_from_blank_and_nothing_is_imported(ctx):
    test_id = make_test(ctx)
    override = {"101005": {1: "AB", **{q: ch for q, ch in enumerate("BCDABCDAB", start=2)}},
                "101003": "----------"}
    response = post_omr(ctx["admin"], test_id, sheets_for(overrides=override), dry_run=False)
    body = response.json()

    assert response.status_code == 200
    assert body["status"] == "review_required"
    assert body["committed"] is False and "Review required" in body["message"]
    assert body["counts"]["multi_mark"] == 1
    assert body["counts"]["blank"] == 12          # blank is counted separately
    assert body["review_items"] == [{
        "sheet": "scan_101005.png", "roll_number": "101005", "question_number": 1,
        "status": "multi_mark", "raw_value": "AB", "code": "answer_multi_mark",
    }]
    assert body["import"] is None
    # the ambiguous answer was NOT quietly stored as a blank
    assert answer_count(test_id) == 0


def test_resolved_sheet_then_imports_cleanly(ctx):
    """The same batch with the multi-mark corrected on the sheet."""
    test_id = make_test(ctx)
    body = post_omr(ctx["admin"], test_id, sheets_for(), dry_run=False).json()
    assert body["status"] == "imported"
    assert student_results(ctx["admin"], test_id)["101005"] == (40, 10, 0, 0)


def test_normalised_omr_answers_evaluate_exactly_like_manual_answers(ctx):
    omr_test = make_test(ctx)
    csv_test = make_test(ctx)

    assert post_omr(ctx["admin"], omr_test, sheets_for(), dry_run=False).json()["status"] == "imported"

    rows = "roll_number,question_number,answer\n" + "".join(
        f"{roll},{q},{'' if ch == '-' else ch}\n"
        for roll, letters in STUDENTS.items()
        for q, ch in enumerate(letters, start=1)
    )
    assert upload(ctx["admin"], "answers", rows, dict(test_id=csv_test, dry_run=False))["status"] == "imported"

    assert student_results(ctx["admin"], omr_test) == student_results(ctx["admin"], csv_test) == EXPECTED


def test_omr_score_never_reaches_the_results(ctx):
    """The recogniser's own score column is ignored; marks come from evaluation."""
    test_id = make_test(ctx)
    post_omr(ctx["admin"], test_id, sheets_for(), dry_run=False)
    marks = {m for m, *_ in student_results(ctx["admin"], test_id).values()}
    assert marks == {40, 22, 0, 10}

    db = SessionLocal()
    try:
        values = {a.answer for a in db.query(models.StudentAnswer).filter_by(test_id=test_id)}
    finally:
        db.close()
    assert values <= {None, "A", "B", "C", "D"}


# ------------------------------------------------------------------
# Atomic, all-or-nothing batches
# ------------------------------------------------------------------

def test_one_bad_sheet_blocks_the_whole_batch(ctx):
    test_id = make_test(ctx)
    files = sheets_for(["101001", "101002", "101003"])
    files.append(("stranger.png", render_sheet("777777", {1: "A"}, seed=9)))

    body = post_omr(ctx["admin"], test_id, files, dry_run=False).json()

    assert body["status"] == "rejected" and body["committed"] is False
    assert [e["code"] for e in body["errors"]] == ["roll_unknown"]
    assert body["counts"]["unknown_student"] == 1 and body["counts"]["accepted"] == 3
    # the three good sheets were not imported either
    assert answer_count(test_id) == 0


def test_duplicate_students_are_reported_not_overwritten(ctx):
    test_id = make_test(ctx)
    files = sheets_for(["101001"]) + [("again.png", render_sheet("101001", {1: "B"}, seed=4))]
    body = post_omr(ctx["admin"], test_id, files, dry_run=False).json()

    assert body["status"] == "rejected"
    assert [e["code"] for e in body["errors"]] == ["roll_duplicate"]
    assert body["counts"]["duplicate"] == 1
    assert answer_count(test_id) == 0


def test_missing_roll_number_is_rejected(ctx):
    test_id = make_test(ctx)
    from omr_sheets import render_sheet as draw
    import numpy as np, cv2
    # a sheet with a roll column left unmarked: 5 digits instead of 6
    png = draw("101001", {1: "A"}, seed=1)
    image = cv2.imdecode(np.frombuffer(png, np.uint8), cv2.IMREAD_GRAYSCALE)
    layout = json.loads(checker_template().layout_path.read_text())
    ox, oy = layout["fieldBlocks"]["Roll"]["origin"]
    cv2.rectangle(image, (ox + 5 * 50 - 6, oy - 6), (ox + 5 * 50 + 36, oy + 10 * 36 + 6), 245, -1)
    ok, blanked = cv2.imencode(".png", image)

    body = post_omr(ctx["admin"], test_id, [("noroll.png", blanked.tobytes())]).json()
    assert body["status"] == "rejected"
    assert body["errors"][0]["code"] in ("roll_malformed", "roll_missing")


def checker_template():
    from backend.app.omr.template import get_template
    return get_template(None)


def test_existing_answers_block_the_import_unless_replace_is_chosen(ctx):
    test_id = make_test(ctx)
    assert post_omr(ctx["admin"], test_id, sheets_for(), dry_run=False).json()["status"] == "imported"

    again = post_omr(ctx["admin"], test_id, sheets_for(), dry_run=False).json()
    assert again["status"] == "rejected"
    assert again["errors"][0]["code"] == "import_rejected"
    assert answer_count(test_id) == 50


# ------------------------------------------------------------------
# Test-specific validation
# ------------------------------------------------------------------

def test_questions_the_test_does_not_have_are_ignored_and_reported(ctx):
    test_id = make_test(ctx)   # a 10-question test; the sheet template has 60
    body = post_omr(ctx["admin"], test_id, sheets_for(["101001"])).json()
    assert body["status"] == "accepted"
    assert any(w["code"] == "template_fields_ignored" for w in body["warnings"])


def test_marks_on_questions_outside_the_test_are_not_imported(ctx):
    test_id = make_test(ctx)
    extra = {q: ch for q, ch in enumerate(KEY, start=1)}
    extra.update({11: "A", 12: "B"})    # the test only has Q1-Q10
    body = post_omr(ctx["admin"], test_id,
                    [("s.png", render_sheet("101001", extra, seed=2))], dry_run=False).json()
    assert body["status"] == "imported"
    assert any(w["code"] == "ignored_fields_marked" for w in body["warnings"])
    assert answer_count(test_id) == 10


def test_a_test_with_more_questions_than_the_template_is_rejected(ctx):
    test_id = make_test(ctx, questions=61)
    body = post_omr(ctx["admin"], test_id, sheets_for(["101001"])).json()
    assert body["status"] == "rejected"
    assert body["errors"][0]["code"] == "template_incompatible"


def test_unknown_template_is_rejected(ctx):
    test_id = make_test(ctx)
    body = post_omr(ctx["admin"], test_id, sheets_for(["101001"]), template_id="nope").json()
    assert body["status"] == "rejected" and body["errors"][0]["code"] == "unknown_template"


def test_a_batch_cannot_enter_another_test_by_id(ctx):
    """Answers go only to the test named in the URL."""
    a, b = make_test(ctx), make_test(ctx)
    post_omr(ctx["admin"], a, sheets_for(["101001"]), dry_run=False)
    assert answer_count(a) == 10 and answer_count(b) == 0


# ------------------------------------------------------------------
# Uploads
# ------------------------------------------------------------------

@pytest.mark.parametrize(
    "name,data",
    [
        ("notes.png", b"this is not an image" * 100),
        ("empty.png", b""),
        ("doc.pdf", b"%PDF-1.4\n" + b"x" * 500),
        ("sheet.gif", b"GIF89a" + b"x" * 500),
    ],
)
def test_invalid_images_are_rejected_and_nothing_is_processed(ctx, monkeypatch, name, data):
    test_id = make_test(ctx)
    ran = []
    monkeypatch.setattr(checker, "run_recognition", lambda ws: ran.append(1))
    import backend.app.omr.service as service
    monkeypatch.setattr(service, "run_recognition", lambda ws: ran.append(1))

    good = sheets_for(["101001"])
    body = post_omr(ctx["admin"], test_id, good + [(name, data)], dry_run=False).json()

    assert body["status"] == "rejected"
    assert [e["code"] for e in body["errors"]] == ["invalid_image"]
    assert body["errors"][0]["sheet"] == name
    assert ran == [] and answer_count(test_id) == 0


def test_oversized_image_is_rejected(ctx, monkeypatch):
    test_id = make_test(ctx)
    monkeypatch.setattr(config, "MAX_OMR_IMAGE_BYTES", 50_000)
    body = post_omr(ctx["admin"], test_id, sheets_for(["101001"])).json()
    assert body["status"] == "rejected"
    assert "larger than" in body["errors"][0]["message"]


def test_oversized_request_is_refused_before_it_is_read(ctx, monkeypatch):
    test_id = make_test(ctx)
    monkeypatch.setattr(config, "MAX_OMR_REQUEST_BYTES", 100_000)
    response = post_omr(ctx["admin"], test_id, sheets_for(["101001"]))
    assert response.status_code == 413
    assert "OMR uploads may be up to" in response.json()["detail"]


def test_too_many_files_are_refused(ctx, monkeypatch):
    test_id = make_test(ctx)
    monkeypatch.setattr(config, "MAX_OMR_FILES", 2)
    body = post_omr(ctx["admin"], test_id, sheets_for(["101001", "101002", "101003"])).json()
    assert body["status"] == "rejected" and body["errors"][0]["code"] == "too_many_files"


def test_no_files_is_a_clear_error(ctx):
    test_id = make_test(ctx)
    body = ctx["admin"].post(f"/api/tests/{test_id}/omr/import", data={}).json()
    assert body["status"] == "rejected" and body["errors"][0]["code"] == "no_files"


# ------------------------------------------------------------------
# Security of processing and responses
# ------------------------------------------------------------------

def test_hostile_filenames_are_never_used_as_paths(ctx, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    test_id = make_test(ctx)
    hostile = {"101001": "../../../../tmp/pwned.png",
               "101002": "..\\..\\windows\\evil.png",
               "101003": "=HYPERLINK(\"http://evil\",\"x\").png"}
    body = post_omr(ctx["admin"], test_id, sheets_for(["101001", "101002", "101003"], names=hostile)).json()

    assert body["status"] == "accepted"
    assert not Path("/tmp/pwned.png").exists()
    labels = [s["sheet"] for s in body["sheets"]]
    assert labels[0] == "pwned.png" and labels[1] == "evil.png"      # directories removed
    assert "/" not in "".join(labels) and "\\" not in "".join(labels)


def test_responses_contain_no_paths_or_engine_internals(ctx, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    test_id = make_test(ctx)
    texts = [
        post_omr(ctx["admin"], test_id, sheets_for(), dry_run=False).text,
        post_omr(ctx["admin"], test_id, sheets_for(["101001"]) + [("x.png", b"junk")]).text,
    ]
    for text in texts:
        for forbidden in (str(tmp_path), "/tmp", "omr-", "third_party", "Traceback",
                          "sheet_0", "input_path", "output_path", "Results_", "score"):
            assert forbidden not in text, forbidden


def test_engine_failure_gives_a_fixed_safe_message_and_cleans_up(ctx, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    test_id = make_test(ctx)

    def boom(*args, **kwargs):
        raise subprocess.SubprocessError(f"secret detail {tmp_path}")

    def exit1(*args, **kwargs):
        return subprocess.CompletedProcess(args, 1, "", f"Traceback: boom in {tmp_path}")

    monkeypatch.setattr(subprocess, "run", exit1)
    response = post_omr(ctx["admin"], test_id, sheets_for(["101001"]), dry_run=False)
    assert response.status_code == 502
    assert response.json() == {"detail": "OMR processing failed."}
    assert "Traceback" not in response.text and str(tmp_path) not in response.text
    assert list((tmp_path / "uploads").iterdir()) == []      # working folder removed
    assert answer_count(test_id) == 0


def test_timeout_is_reported_safely(ctx, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    test_id = make_test(ctx)

    def slow(*args, **kwargs):
        raise subprocess.TimeoutExpired(args, 1)

    monkeypatch.setattr(subprocess, "run", slow)
    response = post_omr(ctx["admin"], test_id, sheets_for(["101001"]))
    assert response.status_code == 502 and "took too long" in response.json()["detail"]
    assert list((tmp_path / "uploads").iterdir()) == []


def test_missing_engine_is_a_clean_503(ctx, monkeypatch):
    test_id = make_test(ctx)
    monkeypatch.setattr(checker, "engine_available", lambda: False)
    response = post_omr(ctx["admin"], test_id, sheets_for(["101001"]))
    assert response.status_code == 503 and "not installed" in response.json()["detail"]


def test_recogniser_is_run_safely_and_headless(ctx, monkeypatch):
    test_id = make_test(ctx)
    seen = {}
    real_run = subprocess.run

    def spy(command, **kwargs):
        seen.update(command=command, **kwargs)
        return real_run(command, **kwargs)

    monkeypatch.setattr(subprocess, "run", spy)
    monkeypatch.setenv("DISPLAY", ":99")
    monkeypatch.setenv("COACHING_SECRET_THING", "do-not-leak")
    post_omr(ctx["admin"], test_id, sheets_for(["101001"], names={"101001": "evil; rm -rf $HOME.png"}))

    assert isinstance(seen["command"], list) and seen["shell"] is False   # no shell
    assert all("evil" not in part for part in seen["command"])            # no user input
    assert seen["cwd"] == str(checker.VENDOR_ROOT)
    assert "DISPLAY" not in seen["env"]                                   # no GUI needed
    assert "COACHING_SECRET_THING" not in seen["env"]
    assert seen["env"]["MPLBACKEND"] == "Agg" and seen["timeout"] > 0


def test_each_batch_gets_a_fresh_output_folder(ctx, tmp_path, monkeypatch):
    """OMRChecker appends to an existing results file, so folders are never shared."""
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    test_id = make_test(ctx)
    outputs = []
    real = checker.run_recognition

    def record(workspace):
        outputs.append(workspace.output_dir)
        assert not workspace.output_dir.exists()          # nothing left from earlier
        return real(workspace)

    import backend.app.omr.service as service
    monkeypatch.setattr(service, "run_recognition", record)

    first = post_omr(ctx["admin"], test_id, sheets_for(["101001", "101002"])).json()
    second = post_omr(ctx["admin"], test_id, sheets_for(["101001", "101002"])).json()

    assert len(set(outputs)) == 2 and not any(p.exists() for p in outputs)
    # a shared results file would have produced 4 rows the second time
    assert first["counts"]["processed"] == second["counts"]["processed"] == 2
    assert second["counts"]["duplicate"] == 0


def test_temporary_files_are_removed_after_success_and_failure(ctx, tmp_path, monkeypatch):
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    test_id = make_test(ctx)

    post_omr(ctx["admin"], test_id, sheets_for(["101001"]))                      # success
    post_omr(ctx["admin"], test_id, sheets_for(["101001"]) + [("x.png", b"j")])  # rejected early
    post_omr(ctx["admin"], test_id, sheets_for(["101001"]),
             template_id="nope")                                                  # rejected early

    assert list((tmp_path / "uploads").iterdir()) == []


# ------------------------------------------------------------------
# Excel and permanent data stay safe
# ------------------------------------------------------------------

def test_formula_like_text_cannot_become_an_excel_formula(ctx):
    """A student called =EVIL()+1 and a hostile filename, after an OMR import."""
    test_id = make_test(ctx)
    names = {"101004": "=cmd|' /C calc'!A0.png"}
    body = post_omr(ctx["admin"], test_id, sheets_for(["101004", "101001"], names=names),
                    dry_run=False).json()
    assert body["status"] == "imported"

    workbook = load_workbook(io.BytesIO(
        ctx["admin"].get(f"/api/tests/{test_id}/export.xlsx").content))
    for sheet in workbook:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (sheet.title, cell.coordinate)

    students = {r[0].value: r[1].value for r in workbook["Student Results"].iter_rows(min_row=2)}
    assert students["101004"] == "'=EVIL()+1"          # neutralised, still readable

    # the upload's filename was never stored anywhere in the workbook
    text = " ".join(str(c.value) for s in workbook for r in s.iter_rows() for c in r if c.value)
    assert "cmd|" not in text and "scan_" not in text


def test_excel_export_still_works_after_an_omr_import(ctx):
    test_id = make_test(ctx)
    post_omr(ctx["admin"], test_id, sheets_for(), dry_run=False)
    response = ctx["admin"].get(f"/api/tests/{test_id}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    rows = {r[0].value: [c.value for c in r] for r in workbook["Student Results"].iter_rows(min_row=2)}
    assert rows["101002"][2:7] == [22, rows["101002"][3], 6, 2, 2]


def test_busy_limit_is_reported_and_the_slot_is_always_released(ctx, monkeypatch):
    import threading
    test_id = make_test(ctx)
    slots = threading.BoundedSemaphore(1)
    monkeypatch.setattr(checker, "_slots", slots)

    slots.acquire()                                    # another batch is running
    busy = post_omr(ctx["admin"], test_id, sheets_for(["101001"]))
    assert busy.status_code == 429 and "busy" in busy.json()["detail"]
    slots.release()

    # a failing run must give the slot back
    monkeypatch.setattr(subprocess, "run",
                        lambda *a, **k: subprocess.CompletedProcess(a, 1, "", "boom"))
    assert post_omr(ctx["admin"], test_id, sheets_for(["101001"])).status_code == 502
    monkeypatch.undo()
    monkeypatch.setattr(checker, "_slots", slots)

    ok = post_omr(ctx["admin"], test_id, sheets_for(["101001"]))
    assert ok.status_code == 200 and ok.json()["status"] == "accepted"


def _corrupt_png():
    """A valid PNG header (so it passes the upload check) with a damaged body."""
    import struct, zlib
    from backend.app.omr.images import PNG_SIGNATURE

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    return (PNG_SIGNATURE
            + chunk(b"IHDR", struct.pack(">IIBBBBB", 1100, 1800, 8, 0, 0, 0, 0))
            + chunk(b"IDAT", b"garbage-not-zlib") + chunk(b"IEND", b""))


def test_an_unreadable_image_is_named_and_does_not_break_the_batch(ctx):
    """Valid header, damaged pixels: the bad file is reported by name."""
    test_id = make_test(ctx)
    files = sheets_for(["101001", "101002"]) + [("damaged.png", _corrupt_png())]

    response = post_omr(ctx["admin"], test_id, files, dry_run=False)
    body = response.json()

    assert response.status_code == 200
    assert body["status"] == "rejected" and body["committed"] is False
    assert [(e["code"], e["sheet"]) for e in body["errors"]] == [("invalid_image", "damaged.png")]
    assert answer_count(test_id) == 0            # the two good sheets were not imported
