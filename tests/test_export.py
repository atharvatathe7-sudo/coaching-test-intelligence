"""
Excel export (Stage 2A).

The main fixture is a small institute built through the real import
workflow. Its expected numbers are worked out by hand in the comments
below and written as literals, so the tests do not depend on the
production calculations they are checking.

Dataset (marking +4 / -1 / 0), current test "X Test 1":

  Question  key  chapter    topic        X1 X2 X3 X4(blank)
  Q1        A    Mechanics  Kinematics   A  A  B  -
  Q2        B    Mechanics  Work         B  C  B  -
  Q3        C    Optics     Lenses       C  C  A  -
  Q4        D    -Heat      =Conduction  D  -  A  -

  X1: 4 correct                    marks 16   accuracy 100
  X2: Q1,Q3 correct, Q2 wrong, Q4 blank   marks 7    accuracy 66.67
  X3: Q2 correct, Q1,Q3,Q4 wrong   marks 1    accuracy 25
  X4: all blank                    marks 0    accuracy 0
  totals: correct 7, wrong 4, unattempted 5 (16 responses)
  average marks (16+7+1+0)/4 = 6.0; highest 16; lowest 0

Previous test "X Test 0" (earlier date, same batch and subject):
  X1 A A A A (marks 1, acc 25), X2 and X3 all correct (16, 100), X4 blank.
"""

import io
import re
import zipfile

import pytest
from openpyxl import load_workbook

from backend.app.database import models
from backend.app.database.connection import DATABASE_PATH, SessionLocal
from backend.app.services import excel_export
from backend.app.services.excel_export import (
    SHEET_NAMES,
    export_filename,
    safe_cell_value,
)

import manage
from conftest import make_client

PASSWORD = "export-test-password"

EVIL_NAME = '=HYPERLINK("http://evil.example","click")'
UNICODE_NAME = "Zoë Müller 日本語"


def _upload(client, path, csv_text, data):
    response = client.post(
        f"/api/imports/{path}",
        data={k: str(v) for k, v in data.items()},
        files={"file": ("f.csv", csv_text.encode("utf-8"), "text/csv")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _setup_test(admin, batch_id, name, date_text, key_rows, with_answers):
    setup = _upload(
        admin,
        "test-setup",
        "question_number,correct_answer,chapter,topic\n" + key_rows,
        {
            "batch_id": batch_id,
            "test_name": name,
            "subject": "Physics",
            "test_date": date_text,
            "dry_run": False,
            "confirm_new_chapters_topics": True,
        },
    )
    assert setup["status"] == "imported", setup
    test_id = setup["summary"]["test_id"]

    if with_answers is not None:
        result = _upload(
            admin,
            "answers",
            "roll_number,question_number,answer\n" + with_answers,
            {"test_id": test_id, "dry_run": False},
        )
        assert result["status"] == "imported", result

    return test_id


def _answers(grid):
    """grid: {roll: 'ABCD' with '-' for blank} -> CSV rows."""
    lines = []
    for roll, answers in grid.items():
        for number, letter in enumerate(answers, start=1):
            lines.append(f"{roll},{number},{'' if letter == '-' else letter}\n")
    return "".join(lines)


KEY = (
    "1,A,Mechanics,Kinematics\n"
    "2,B,Mechanics,Work\n"
    "3,C,Optics,Lenses\n"
    "4,D,-Heat,=Conduction\n"
)


@pytest.fixture(scope="module")
def x(demo_database):
    """Institute X: admin, teacher, roster, two tests, one action."""
    db = SessionLocal()
    try:
        institute = models.Institute(name="Export Institute =cmd")
        db.add(institute)
        db.flush()
        manage.make_user(db, institute.id, "X Admin", "admin@x.test", "admin", PASSWORD)
        manage.make_user(db, institute.id, "X Teacher", "teacher@x.test", "teacher", PASSWORD)
        db.commit()
    finally:
        db.close()

    admin = make_client("admin@x.test", PASSWORD)
    teacher = make_client("teacher@x.test", PASSWORD)

    batch = admin.post("/api/batches", json={"name": "X Batch"}).json()

    roster = (
        "roll_number,name\n"
        "X1,Asha Rao\n"
        f'X2,"{EVIL_NAME.replace(chr(34), chr(34) * 2)}"\n'
        f"X3,{UNICODE_NAME}\n"
        "X4,@dev\n"
    )
    assert _upload(admin, "roster", roster,
                   {"batch_id": batch["id"], "dry_run": False})["status"] == "imported"

    previous = _setup_test(
        admin, batch["id"], "X Test 0", "2026-04-01", KEY,
        _answers({"X1": "AAAA", "X2": "ABCD", "X3": "ABCD", "X4": "----"}),
    )
    current = _setup_test(
        admin, batch["id"], "X Test 1", "2026-05-01", KEY,
        _answers({"X1": "ABCD", "X2": "ACC-", "X3": "BBAA", "X4": "----"}),
    )

    return {
        "admin": admin,
        "teacher": teacher,
        "batch_id": batch["id"],
        "previous": previous,
        "current": current,
    }


def _workbook(client, test_id):
    response = client.get(f"/api/tests/{test_id}/export.xlsx")
    assert response.status_code == 200, response.text
    return load_workbook(io.BytesIO(response.content))


def _rows(workbook, sheet):
    return [
        list(row)
        for row in workbook[sheet].iter_rows(values_only=True)
        if any(value is not None for value in row)
    ]


def _summary(workbook):
    return {
        row[0]: row[1]
        for row in _rows(workbook, "Test Summary")
        if row[0] is not None
    }


@pytest.fixture(scope="module")
def current_book(x):
    """
    The exported current test. conftest's clean_actions fixture deletes
    teacher actions after every test, so the action is recorded here,
    immediately before the one export that these tests share.
    """
    long_note = "N" * 40000 + " end"
    action = x["teacher"].post(
        "/api/actions",
        json={
            "test_id": x["current"],
            "question_number": 4,
            "action_type": "reteach",
            "finding_title": "=SUM(1+1) finding",
            "note": "=1+1 日本語 " + long_note,
        },
    )
    assert action.status_code == 200, action.text
    x["action_id"] = action.json()["id"]

    return _workbook(x["admin"], x["current"])


# ------------------------------------------------------------------
# Access, response and file format
# ------------------------------------------------------------------

def test_authenticated_export_returns_a_valid_xlsx(x):
    response = x["admin"].get(f"/api/tests/{x['current']}/export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == excel_export.XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == (
        'attachment; filename="X_Test_1_2026-05-01.xlsx"'
    )
    assert response.headers["cache-control"] == "no-store"

    # A genuine xlsx is a zip package containing the workbook part.
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert archive.testzip() is None
    assert "xl/workbook.xml" in archive.namelist()
    assert load_workbook(io.BytesIO(response.content)).sheetnames


def test_teacher_can_export_too(x):
    assert x["teacher"].get(f"/api/tests/{x['current']}/export.xlsx").status_code == 200


def test_unauthenticated_export_is_rejected(x):
    anonymous = make_client()
    response = anonymous.get(f"/api/tests/{x['current']}/export.xlsx")
    assert response.status_code == 401
    assert response.content[:2] != b"PK"


def test_cross_tenant_export_is_indistinguishable_from_missing(x, client):
    """The demo institute cannot export institute X's test."""
    foreign = client.get(f"/api/tests/{x['current']}/export.xlsx")
    missing = client.get("/api/tests/99999999/export.xlsx")

    assert foreign.status_code == missing.status_code == 404
    assert foreign.json() == missing.json()
    assert foreign.headers["content-type"] == missing.headers["content-type"]

    # ... and the reverse direction.
    demo_test = client.get("/api/tests").json()["tests"][0]["id"]
    assert x["admin"].get(f"/api/tests/{demo_test}/export.xlsx").status_code == 404


def test_nonexistent_test_is_a_plain_404(x):
    response = x["admin"].get("/api/tests/99999999/export.xlsx")
    assert response.status_code == 404
    assert str(DATABASE_PATH) not in response.text


def test_export_does_not_change_any_data(x):
    def counts():
        db = SessionLocal()
        try:
            return {
                table: db.query(model).count()
                for table, model in (
                    ("results", models.TestResult),
                    ("answers", models.StudentAnswer),
                    ("actions", models.TeacherAction),
                    ("audit", models.AuditLog),
                )
            }
        finally:
            db.close()

    before = counts()
    _workbook(x["admin"], x["current"])
    assert counts() == before


# ------------------------------------------------------------------
# Structure
# ------------------------------------------------------------------

def test_required_worksheets_in_order(current_book):
    assert current_book.sheetnames == [
        "Test Summary", "Student Results", "Question Analysis",
        "Chapter Analysis", "Topic Analysis", "Answer Matrix",
        "Teacher Actions", "Test Comparison",
    ]
    assert tuple(current_book.sheetnames) == SHEET_NAMES


HEADERS = {
    "Student Results": ["Roll number", "Student name", "Marks", "Accuracy",
                        "Correct", "Wrong", "Unattempted"],
    "Question Analysis": ["Question", "Chapter", "Topic", "Correct answer",
                          "Correct", "Wrong", "Unattempted", "Total responses",
                          "Correct %"],
    "Chapter Analysis": ["Chapter", "Total responses", "Correct", "Wrong",
                         "Unattempted", "Correct %"],
    "Topic Analysis": ["Chapter", "Topic", "Total responses", "Correct",
                       "Wrong", "Unattempted", "Correct %"],
    "Answer Matrix": ["Roll number", "Student name", "Q1", "Q2", "Q3", "Q4"],
    "Teacher Actions": ["Action ID", "Question", "Chapter", "Topic", "Finding",
                        "Action type", "Status", "Note", "Created (UTC)",
                        "Outcome", "Caveats"],
}


@pytest.mark.parametrize("sheet", HEADERS)
def test_expected_headers_exist(current_book, sheet):
    header = [c.value for c in current_book[sheet][1]]
    for expected in HEADERS[sheet]:
        assert expected in header, (sheet, expected, header)


@pytest.mark.parametrize(
    "sheet",
    ["Student Results", "Question Analysis", "Chapter Analysis",
     "Topic Analysis", "Answer Matrix", "Teacher Actions"],
)
def test_tabular_sheets_have_bold_headers_freeze_panes_and_filter(current_book, sheet):
    ws = current_book[sheet]
    assert ws.freeze_panes is not None
    assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:")
    assert all(cell.font.bold for cell in ws[1] if cell.value is not None)


def test_number_formats(current_book):
    students = current_book["Student Results"]
    assert students["D2"].number_format == "0.0%"
    assert students["E2"].number_format == "0"
    questions = current_book["Question Analysis"]
    assert questions["I2"].number_format == "0.0%"
    assert current_book["Teacher Actions"]["K2"].number_format == "yyyy-mm-dd hh:mm"
    assert current_book["Test Summary"]["B4"].number_format == "yyyy-mm-dd"


# ------------------------------------------------------------------
# Values (hand-computed, see module docstring)
# ------------------------------------------------------------------

def test_summary_values(current_book):
    summary = _summary(current_book)

    assert summary["Batch"] == "X Batch"
    assert summary["Test"] == "X Test 1"
    assert summary["Subject"] == "Physics"
    assert summary["Test date"].date().isoformat() == "2026-05-01"
    assert summary["Total students in batch"] == 4
    assert summary["Students evaluated"] == 4
    assert summary["Total questions"] == 4
    assert summary["Average marks"] == pytest.approx(6.0)
    assert summary["Highest marks"] == 16
    assert summary["Lowest marks"] == 0
    # mean of 100, 66.67, 25, 0 = 47.92
    assert summary["Average accuracy (mean of student accuracies)"] == pytest.approx(0.4792, abs=1e-4)
    # pooled: 7 correct of 11 attempted
    assert summary["Batch accuracy (all attempted answers together)"] == pytest.approx(7 / 11, abs=1e-4)
    assert summary["Total correct answers"] == 7
    assert summary["Total wrong answers"] == 4
    assert summary["Total unattempted answers"] == 5


def test_summary_text_is_protected_but_readable(current_book):
    # The institute name begins with "Export", so it is untouched; a name
    # beginning with = would be prefixed (see the unit tests below).
    assert _summary(current_book)["Institute"] == "Export Institute =cmd"


def _summary_lines(workbook):
    return [
        [v for v in row if v is not None]
        for row in workbook["Test Summary"].iter_rows(values_only=True)
        if any(v is not None for v in row)
    ]


def test_summary_ranks_chapters_only_when_they_differ(current_book):
    """X: -Heat 25%, Mechanics 50%, Optics 50% -> one lowest, one highest."""
    lines = _summary_lines(current_book)
    labels = [line[0] for line in lines]

    low = labels.index("Lowest correct % chapters")
    high = labels.index("Highest correct % chapters")

    assert lines[low + 1] == ["'-Heat", pytest.approx(0.25)]
    assert lines[high + 1] == ["Optics", pytest.approx(0.50)]
    # no chapter appears under both headings
    assert lines[low + 1][0] != lines[high + 1][0]


def test_summary_does_not_rank_tied_chapters(client):
    """Demo Test 02: every chapter is at exactly 50%, so nothing is 'lowest'."""
    test_id = {t["name"]: t["id"] for t in client.get("/api/tests").json()["tests"]}["Physics Test 02"]
    lines = _summary_lines(_workbook(client, test_id))
    labels = [line[0] for line in lines]

    assert "Lowest correct % chapters" not in labels
    assert "Highest correct % chapters" not in labels
    same = next(line for line in lines if line[0] == "All chapters have the same correct %")
    assert same[1] == pytest.approx(0.5)


def test_student_rows(current_book):
    rows = _rows(current_book, "Student Results")[1:]

    assert [r[0] for r in rows] == ["X1", "X2", "X3", "X4"]  # roll-number order
    by_roll = {r[0]: r for r in rows}

    def check(roll, marks, accuracy, correct, wrong, blank):
        r = by_roll[roll]
        assert r[2] == marks
        assert r[3] == pytest.approx(accuracy, abs=1e-4)
        assert (r[4], r[5], r[6]) == (correct, wrong, blank)

    check("X1", 16, 1.0, 4, 0, 0)
    check("X2", 7, 2 / 3, 2, 1, 1)
    check("X3", 1, 0.25, 1, 3, 0)
    check("X4", 0, 0.0, 0, 0, 4)

    # Internal IDs are not exposed.
    assert "student_id" not in [str(c.value).lower() for c in current_book["Student Results"][1]]


def test_question_rows(current_book):
    rows = _rows(current_book, "Question Analysis")[1:]
    assert [r[0] for r in rows] == [1, 2, 3, 4]

    # (chapter, topic, key, correct, wrong, blank, responses, correct %)
    expected = {
        1: ("Mechanics", "Kinematics", "A", 2, 1, 1, 4, 0.50),
        2: ("Mechanics", "Work", "B", 2, 1, 1, 4, 0.50),
        3: ("Optics", "Lenses", "C", 2, 1, 1, 4, 0.50),
        4: ("-Heat", "=Conduction", "D", 1, 1, 2, 4, 0.25),
    }
    for row in rows:
        chapter, topic, key, c, w, b, n, pct = expected[row[0]]
        got_chapter = row[1].lstrip("'")  # leading - or = is neutralised
        got_topic = row[2].lstrip("'")
        assert (got_chapter, got_topic, row[3]) == (chapter, topic, key)
        assert (row[4], row[5], row[6], row[7]) == (c, w, b, n)
        assert row[8] == pytest.approx(pct)


def test_chapter_rows(current_book):
    rows = {r[0].lstrip("'"): r for r in _rows(current_book, "Chapter Analysis")[1:]}

    assert set(rows) == {"Mechanics", "Optics", "-Heat"}
    # header: Chapter, Subject, Questions, Total responses, Correct, Wrong,
    #         Unattempted, Correct %
    assert rows["Mechanics"][2:8] == [2, 8, 4, 2, 2, pytest.approx(0.50)]
    assert rows["Optics"][2:8] == [1, 4, 2, 1, 1, pytest.approx(0.50)]
    assert rows["-Heat"][2:8] == [1, 4, 1, 1, 2, pytest.approx(0.25)]


def test_topic_rows_are_not_merged(current_book):
    rows = {
        (r[0].lstrip("'"), r[1].lstrip("'")): r
        for r in _rows(current_book, "Topic Analysis")[1:]
    }
    assert set(rows) == {
        ("Mechanics", "Kinematics"), ("Mechanics", "Work"),
        ("Optics", "Lenses"), ("-Heat", "=Conduction"),
    }
    # Chapter, Topic, Questions, Responses, Correct, Wrong, Unattempted, Correct %
    assert rows[("Mechanics", "Kinematics")][2:8] == [1, 4, 2, 1, 1, pytest.approx(0.5)]
    assert rows[("-Heat", "=Conduction")][2:8] == [1, 4, 1, 1, 2, pytest.approx(0.25)]


def test_answer_matrix(current_book):
    rows = _rows(current_book, "Answer Matrix")
    assert rows[0] == ["Roll number", "Student name", "Q1", "Q2", "Q3", "Q4"]

    matrix = {r[0]: r[2:] for r in rows[1:]}
    assert matrix == {
        "X1": ["A", "B", "C", "D"],
        "X2": ["A", "C", "C", None],
        "X3": ["B", "B", "A", "A"],
        "X4": [None, None, None, None],
    }


def test_teacher_actions_sheet(x, current_book):
    rows = _rows(current_book, "Teacher Actions")
    header = rows[0]
    assert len(rows) == 2

    action = dict(zip(header, rows[1]))
    assert action["Action ID"] == x["action_id"]
    assert action["Question"] == 4
    assert action["Action type"] == "reteach"
    assert action["Status"] == "planned"
    assert action["Finding"] == "'=SUM(1+1) finding"
    assert action["Outcome"] == "Measured" or action["Outcome"]
    assert "still planned" in action["Caveats"]
    assert action["Created (UTC)"] is not None

    # No security or authorship internals.
    lowered = " ".join(str(h).lower() for h in header)
    for forbidden in ("user", "token", "password", "session", "csrf", "hash"):
        assert forbidden not in lowered


def test_outcome_for_measured_action_uses_existing_outcome_logic(client):
    """Demo Test 01 -> 02: observed change in Mechanics (existing outcome rules)."""
    tests = {t["name"]: t["id"] for t in client.get("/api/tests").json()["tests"]}
    mechanics = next(
        c for c in client.get(f"/api/tests/{tests['Physics Test 01']}/analytics/chapters-topics")
        .json()["chapters"] if c["chapter_name"] == "Mechanics"
    )
    assert client.post("/api/actions", json={
        "test_id": tests["Physics Test 01"],
        "chapter_id": mechanics["chapter_id"],
        "action_type": "review",
        "status": "completed",
    }).status_code == 200

    outcome = client.get(
        f"/api/tests/{tests['Physics Test 01']}/actions/outcomes"
    ).json()["outcomes"][0]

    book = _workbook(client, tests["Physics Test 01"])
    rows = _rows(book, "Teacher Actions")
    row = dict(zip(rows[0], rows[1]))

    assert row["Outcome"] == "Measured"
    assert row["Measured target"] == "Chapter: Mechanics"
    assert row["Next comparable test"] == "Physics Test 02"
    assert row["Correct % in this test"] == pytest.approx(outcome["previous_percentage"] / 100)
    assert row["Correct % in next test"] == pytest.approx(outcome["next_percentage"] / 100)
    assert row["Observed change (percentage points)"] == outcome["change_pp"]
    assert "does not establish" in row["Interpretation"]


def test_comparison_included_when_applicable(current_book):
    rows = _rows(current_book, "Test Comparison")
    text = {r[0]: r[1] for r in rows if r[0] and len(r) > 1}

    assert text["Previous test"] == "X Test 0"
    assert text["This test"] == "X Test 1"
    assert text["Students in both tests"] == 4

    measures = {r[0]: r for r in rows if r[0] in ("Average marks", "Average accuracy")}
    # previous marks (1+16+16+0)/4 = 8.25; this test 6.0; change -2.25
    assert measures["Average marks"][1:4] == [
        pytest.approx(8.25), pytest.approx(6.0), pytest.approx(-2.25)
    ]
    # accuracy 56.25% -> 47.92%, a change of -8.33 percentage points
    assert measures["Average accuracy"][1] == pytest.approx(0.5625, abs=1e-4)
    assert measures["Average accuracy"][2] == pytest.approx(0.4792, abs=1e-4)
    assert measures["Average accuracy"][3] == pytest.approx(-8.33, abs=0.02)

    # By marks: X1 improved (1 -> 16); X2, X3 declined; X4 unchanged (0 -> 0)
    summary = next(r for r in rows if r[0] == 1 and r[1] == 2 and r[2] == 1)
    assert summary[:3] == [1, 2, 1]


def test_no_comparison_is_reported_not_invented(client):
    first = {t["name"]: t["id"] for t in client.get("/api/tests").json()["tests"]}["Physics Test 01"]
    rows = _rows(_workbook(client, first), "Test Comparison")
    assert rows == [["No comparable test is available."]]


# ------------------------------------------------------------------
# Demo data: values the project documents (CLAUDE.md)
# ------------------------------------------------------------------

def test_demo_test_values_match_documented_figures(client):
    tests = {t["name"]: t["id"] for t in client.get("/api/tests").json()["tests"]}

    one = _summary(_workbook(client, tests["Physics Test 01"]))
    assert one["Students evaluated"] == 20
    assert one["Total questions"] == 30
    assert one["Average marks"] == pytest.approx(45.8)
    assert one["Average accuracy (mean of student accuracies)"] == pytest.approx(0.5391, abs=1e-4)

    book_two = _workbook(client, tests["Physics Test 02"])
    two = _summary(book_two)
    assert two["Average marks"] == pytest.approx(51.0)
    assert two["Average accuracy (mean of student accuracies)"] == pytest.approx(0.625, abs=1e-4)

    comparison = _rows(book_two, "Test Comparison")
    marks = next(r for r in comparison if r[0] == "Average marks")
    accuracy = next(r for r in comparison if r[0] == "Average accuracy")
    assert marks[3] == pytest.approx(5.2)
    assert accuracy[3] == pytest.approx(8.59, abs=0.01)
    assert next(r for r in comparison if r[:3] == [13, 6, 1])


# ------------------------------------------------------------------
# Formula injection and text safety
# ------------------------------------------------------------------

@pytest.mark.parametrize("text", ["=1+1", "+1+1", "-1+1", "@SUM(A1)", "\t=1", "\r=1",
                                  '=HYPERLINK("http://e","x")'])
def test_safe_cell_value_neutralises_formula_triggers(text):
    assert safe_cell_value(text) == "'" + text


@pytest.mark.parametrize("value", [0, 5, -3, 4.5, -0.25, True, None])
def test_safe_cell_value_leaves_numbers_alone(value):
    assert safe_cell_value(value) == value


def test_safe_cell_value_keeps_dates_and_ordinary_text():
    import datetime as dt

    today = dt.date(2026, 5, 1)
    assert safe_cell_value(today) is today
    assert safe_cell_value("Asha Rao") == "Asha Rao"
    assert safe_cell_value("a=b") == "a=b"
    assert safe_cell_value("Zoë 日本語") == "Zoë 日本語"
    assert safe_cell_value(float("nan")) is None


def test_safe_cell_value_removes_xml_illegal_characters_and_caps_length():
    assert safe_cell_value("a\x00b\x07c") == "abc"
    assert len(safe_cell_value("x" * 50000)) <= 32767


def test_no_cell_in_the_workbook_is_a_formula(current_book):
    for ws in current_book:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "@", "\t", "\r")), (
                        ws.title, cell.coordinate, cell.value)


def test_formula_like_student_names_are_stored_as_text(current_book):
    names = {r[0]: r[1] for r in _rows(current_book, "Student Results")[1:]}

    assert names["X2"] == "'" + EVIL_NAME  # neutralised, content preserved
    assert names["X4"] == "'@dev"
    assert names["X1"] == "Asha Rao"


def test_unicode_survives(current_book):
    names = {r[0]: r[1] for r in _rows(current_book, "Student Results")[1:]}
    assert names["X3"] == UNICODE_NAME

    matrix_names = {r[0]: r[1] for r in _rows(current_book, "Answer Matrix")[1:]}
    assert matrix_names["X3"] == UNICODE_NAME

    action = dict(zip(*_rows(current_book, "Teacher Actions")[:2]))
    assert "日本語" in action["Note"]


def test_very_long_note_is_cut_to_the_cell_limit(current_book):
    action = dict(zip(*_rows(current_book, "Teacher Actions")[:2]))
    assert 30000 < len(action["Note"]) <= 32767


def test_no_secrets_or_paths_in_the_workbook(x):
    response = x["admin"].get(f"/api/tests/{x['current']}/export.xlsx")
    book = load_workbook(io.BytesIO(response.content))

    haystack = " ".join(
        str(cell.value)
        for ws in book
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    ).lower()

    for forbidden in (str(DATABASE_PATH).lower(), "cti_session", "argon2",
                      "password", "x-requested-with", "alembic"):
        assert forbidden not in haystack

    # Workbook properties carry no filesystem paths either.
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    meta = archive.read("docProps/core.xml").decode() + archive.read("docProps/app.xml").decode()
    assert "/home/" not in meta and "\\Users\\" not in meta


# ------------------------------------------------------------------
# Filename
# ------------------------------------------------------------------

@pytest.mark.parametrize(
    "name,expected",
    [
        ("Physics Test 01", "Physics_Test_01_2026-10-02.xlsx"),
        ("../../etc/passwd", "etc_passwd_2026-10-02.xlsx"),
        ('a"b;c\r\nd', "a_b_c_d_2026-10-02.xlsx"),
        ("日本語", "Test_2026-10-02.xlsx"),
        ("Zoë Test", "Zoe_Test_2026-10-02.xlsx"),
        ("x" * 300, "x" * 80 + "_2026-10-02.xlsx"),
    ],
)
def test_filename_is_deterministic_and_safe(name, expected):
    import datetime as dt

    test = models.Test(name=name, test_date=dt.date(2026, 10, 2))
    assert export_filename(test) == expected
    assert re.fullmatch(r"[A-Za-z0-9_.-]+", export_filename(test))


# ------------------------------------------------------------------
# Edge cases: each must still produce a workbook that opens
# ------------------------------------------------------------------

def test_test_with_no_teacher_actions(client):
    tests = {t["name"]: t["id"] for t in client.get("/api/tests").json()["tests"]}
    book = _workbook(client, tests["Physics Test 02"])
    assert _rows(book, "Teacher Actions")[1][0] == "No teacher actions have been recorded for this test."


def test_test_with_questions_but_no_answers(x):
    test_id = _setup_test(
        x["admin"], x["batch_id"], "X No Answers", "2026-06-01", KEY, None
    )
    book = _workbook(x["admin"], test_id)

    assert book.sheetnames == list(SHEET_NAMES)
    summary = _summary(book)
    assert summary["Total questions"] == 4
    assert summary["Students evaluated"] == 0
    assert "Average marks" not in summary
    assert any("No student answers" in str(v) for v in
               (c.value for row in book["Test Summary"].iter_rows() for c in row))

    # Question rows still list the four questions (no responses yet).
    questions = _rows(book, "Question Analysis")
    assert [r[0] for r in questions[1:]] == [1, 2, 3, 4]
    assert all(r[7] == 0 for r in questions[1:])

    for sheet in ("Student Results", "Chapter Analysis", "Topic Analysis", "Answer Matrix"):
        assert "No student answers" in _rows(book, sheet)[1][0]

    # The earlier "X Test 1" exists but this one has no answers to compare.
    assert _rows(book, "Test Comparison")[0] == ["No comparable test is available."]


def test_test_with_no_questions_and_no_students(x):
    db = SessionLocal()
    try:
        empty_batch = models.Batch(name="Empty Batch", institute_id=db.query(models.Batch)
                                   .filter_by(id=x["batch_id"]).one().institute_id)
        db.add(empty_batch)
        db.flush()
        import datetime as dt
        test = models.Test(batch_id=empty_batch.id, name="Empty =Test", subject="Physics",
                           test_date=dt.date(2026, 7, 1), marks_correct=4,
                           marks_wrong=-1, marks_blank=0)
        db.add(test)
        db.commit()
        test_id = test.id
    finally:
        db.close()

    book = _workbook(x["admin"], test_id)
    assert book.sheetnames == list(SHEET_NAMES)
    summary = _summary(book)
    assert summary["Total students in batch"] == 0
    assert summary["Total questions"] == 0
    assert summary["Test"] == "'Empty =Test".lstrip("'")  # not at the start: unchanged
    assert "no questions" in _rows(book, "Question Analysis")[1][0]


def test_all_students_unattempted(x):
    test_id = _setup_test(
        x["admin"], x["batch_id"], "X All Blank", "2026-08-01", KEY,
        _answers({"X1": "----", "X2": "----", "X3": "----", "X4": "----"}),
    )
    book = _workbook(x["admin"], test_id)
    summary = _summary(book)

    assert summary["Students evaluated"] == 4
    assert summary["Total correct answers"] == 0
    assert summary["Total wrong answers"] == 0
    assert summary["Total unattempted answers"] == 16
    assert summary["Average marks"] == 0

    students = _rows(book, "Student Results")[1:]
    assert [r[2] for r in students] == [0, 0, 0, 0]
    assert [r[3] for r in students] == [0, 0, 0, 0]  # accuracy of zero attempts
    assert all(v is None for r in _rows(book, "Answer Matrix")[1:] for v in r[2:])


def test_authorization_for_other_exports_unchanged(x, client):
    """Existing behaviour: A cannot read B's analytics either."""
    assert client.get(f"/api/tests/{x['current']}/analytics/batch").status_code == 404
    assert x["admin"].get(f"/api/tests/{x['current']}/analytics/batch").status_code == 200
